import csv
import io
from app.models.audience_member import AudienceMember
from fastapi import APIRouter, Depends, HTTPException, Query, status, File, UploadFile
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.user import User
from app.schemas.audience import (
    AudienceCreate,
    AudienceResponse,
    AudienceUpdate,
)
from app.services import audience_service
from openpyxl import load_workbook
from pydantic import ValidationError

router = APIRouter(
    prefix="/audience",
    tags=["Audience Management"],
)


@router.post(
    "/",
    response_model=AudienceResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_audience(
    data: AudienceCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        return audience_service.create_audience_member(db, data)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="An audience member with this email or phone already exists.",
        )


@router.get("/")
def list_audience(
    skip: int = Query(0, ge=0),
    limit: int = Query(10, ge=1, le=100),
    search: str | None = None,
    language: str | None = None,
    city: str | None = None,
    occupation: str | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return audience_service.get_audience_list(
        db=db,
        skip=skip,
        limit=limit,
        search=search,
        language=language,
        city=city,
        occupation=occupation,
    )


@router.get("/stats")
def audience_stats(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return audience_service.get_audience_stats(db)


@router.get("/{audience_id}", response_model=AudienceResponse)
def get_audience(
    audience_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    member = audience_service.get_audience_member(db, audience_id)

    if not member:
        raise HTTPException(status_code=404, detail="Audience member not found")

    return member


@router.put("/{audience_id}", response_model=AudienceResponse)
def update_audience(
    audience_id: int,
    data: AudienceUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        member = audience_service.update_audience_member(
            db, audience_id, data
        )
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="An audience member with this email or phone already exists.",
        )

    if not member:
        raise HTTPException(status_code=404, detail="Audience member not found")

    return member


@router.delete("/{audience_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_audience(
    audience_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    deleted = audience_service.delete_audience_member(db, audience_id)

    if not deleted:
        raise HTTPException(status_code=404, detail="Audience member not found")


@router.post("/bulk-upload")
async def bulk_upload_audience(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    filename = (file.filename or "").lower()

    if not filename.endswith((".csv", ".xlsx")):
        raise HTTPException(
            status_code=400,
            detail="Only CSV and XLSX files are supported.",
        )

    content = await file.read()

    try:
        if filename.endswith(".csv"):
            text_content = content.decode("utf-8-sig")
            rows = list(csv.DictReader(io.StringIO(text_content)))
        else:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            headers = next(values, None)

            if not headers:
                raise ValueError("The Excel file is empty.")

            headers = [
                str(header).strip() if header is not None else ""
                for header in headers
            ]
            rows = [
                dict(zip(headers, row))
                for row in values
                if any(value is not None for value in row)
            ]
            workbook.close()

    except (UnicodeDecodeError, csv.Error, ValueError, OSError) as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read the file: {exc}",
        )

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="The file contains no audience records.",
        )
    
    added = 0
    failed = []

    from app.services.audience_service import _check_duplicates

    for row_number, row in enumerate(rows, start=2):
        try:
            with db.begin_nested():
                normalized = {
                    str(key).strip(): value
                    for key, value in row.items()
                    if key is not None and str(key).strip()
                }

                text_fields = (
                    "name", "email", "phone", "gender", "city", "state",
                    "language", "occupation", "organization_department",
                    "preferred_channel",
                )

                for field in text_fields:
                    if normalized.get(field) is not None:
                        normalized[field] = (
                            str(normalized[field]).strip() or None
                        )

                if normalized.get("age") not in (None, ""):
                    normalized["age"] = int(normalized["age"])
                elif normalized.get("age") == "":
                    normalized["age"] = None

                if normalized.get("segment_id") in ("", None):
                    normalized["segment_id"] = None
                else:
                    normalized["segment_id"] = int(
                        normalized["segment_id"]
                    )

                consent = normalized.get("consent_status", False)
                if isinstance(consent, str):
                    consent_text = consent.strip().lower()

                    if consent_text in ("true", "1", "yes"):
                        normalized["consent_status"] = True
                    elif consent_text in ("false", "0", "no", ""):
                        normalized["consent_status"] = False
                    else:
                        raise ValueError(
                            "consent_status must be true/false, "
                            "yes/no, or 1/0."
                        )

                data = AudienceCreate(**normalized)

                _check_duplicates(
                    db,
                    email=str(data.email) if data.email else None,
                    phone=data.phone,
                )

                member = AudienceMember(**data.model_dump())
                db.add(member)
                db.flush()

            # Count the row only after its savepoint succeeds.
            added += 1

        except (ValidationError, ValueError, TypeError, IntegrityError) as exc:
            failed.append({
                "row": row_number,
                "reason": str(exc),
            })
        except Exception as exc:
            failed.append({
                "row": row_number,
                "reason": str(exc),
            })

    try:
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="The bulk upload could not be saved. Please retry.",
        )

    return {
        "total_rows": len(rows),
        "added": added,
        "failed_count": len(failed),
        "failed_rows": failed,
    }
